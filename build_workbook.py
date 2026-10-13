"""Build the ACEG coding workbook used in Weeks 2–5 (pilot, double coding, adjudication).

Run: python tools/build_workbook.py
Out: coding/ACEG_coding_workbook.xlsx
"""
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

CASES = Path("data/processed/cases_to_code.csv")
OUT = Path("coding/ACEG_coding_workbook.xlsx")
FONT = "Arial"
YELLOW = PatternFill("solid", fgColor="FFF7CC")
HEAD = PatternFill("solid", fgColor="DBE7F3")
GREY = PatternFill("solid", fgColor="E5E7EB")
CHECK = PatternFill("solid", fgColor="FDE2E2")
thin = Side(style="thin", color="C7CCD3")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

DOMAINS = [
    ("D1", "Agency and action control", "A", [1, 2, 3]),
    ("D2", "Instruction and boundary adherence", "A", [4, 5]),
    ("D3", "Transparency and honest reporting", "D", [6, 7]),
    ("D4", "Input trust and goal integrity", "B", [8, 9]),
    ("D5", "Secrets and data protection", "C", [10, 11]),
    ("D6", "Code and supply-chain integrity", "B", [12, 13, 14]),
    ("D7", "Accountability and recoverability", "D", [15, 16, 17]),
    ("D8", "Human oversight and appropriate reliance", "A", [18, 19, 20]),
]
CONTROLS = {
    1: "Least-privilege credentials and tool permissions", 2: "Human approval gate for sensitive or irreversible actions",
    3: "Execution isolation (sandbox; dev/prod separation)", 4: "Machine-enforced policy limits (deny rules, ignore files, allowlists)",
    5: "Workspace trust and scope confirmation before execution", 6: "Faithful disclosure of planned actions (full arguments shown)",
    7: "Tamper-evident action logs", 8: "Isolation and sanitization of untrusted inputs",
    9: "Allow-listed, integrity-checked tools and MCP servers", 10: "Secret scanning, redaction and protected credential storage",
    11: "Network egress limits and data-handling policy", 12: "Dependency, extension and release integrity verification",
    13: "Security testing gate before release or merge", 14: "No auto-execution of untrusted workspace configuration",
    15: "Named human owner and authorization for agent actions", 16: "Reversible actions (checkpoints, backup and rollback)",
    17: "Incident response, patching and disclosure process", 18: "Autonomy level matched to risk (safe defaults for auto modes)",
    19: "Human review of agent-authored changes and commands", 20: "Capability and limitation disclosure to users",
}
DOM_OF = {c: d for d, _, _, cs in DOMAINS for c in cs}

# Coder-sheet layout
COL_ID, COL_BATCH, COL_PROD, COL_TITLE = 1, 2, 3, 4
COL_D = 5                      # d1..d8 -> columns 5..12
COL_PRIMARY = 13
COL_C = 14                     # c01..c20 -> columns 14..33
COL_COVERED = 34
COL_NDOM, COL_NAPP, COL_NBR, COL_GCS, COL_CHECK, COL_DONE, COL_NOTES = 35, 36, 37, 38, 39, 40, 41
HEADERS = (["case_id", "batch", "product", "title"] + [f"d{i}" for i in range(1, 9)] + ["primary_domain"]
           + [f"c{i:02d}" for i in range(1, 21)] + ["covered", "n_domains", "n_applicable", "n_breached", "gcs",
                                                     "check", "done", "notes"])
L = get_column_letter


def style_header(ws, ncols, row=1):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = Font(name=FONT, bold=True, size=10)
        cell.fill = HEAD
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER


def pilot_sample(d: pd.DataFrame, n=20, seed=42) -> list[str]:
    rng = np.random.default_rng(seed)
    strata = d.groupby(["component_type", "high_impact"])
    quota = {"agent product": {1: 11, 0: 4}, "ecosystem": {1: 3, 0: 2}}
    picks = []
    for (ct, hi), g in strata:
        k = quota[ct][hi]
        picks += list(rng.choice(g["case_id"].values, size=k, replace=False))
    assert len(picks) == n
    return sorted(picks)


def main():
    d = pd.read_csv(CASES)
    pilot = set(pilot_sample(d))
    rest = [c for c in d["case_id"] if c not in pilot]
    batch = {c: "W2 pilot" for c in pilot}
    for i, c in enumerate(rest):
        batch[c] = "W3" if i < 68 else ("W4" if i < 136 else "W5")
    d["batch"] = d["case_id"].map(batch)
    order = {"W2 pilot": 0, "W3": 1, "W4": 2, "W5": 3}
    d = d.sort_values(["batch", "case_id"], key=lambda s: s.map(order) if s.name == "batch" else s).reset_index(drop=True)
    n = len(d)
    last = n + 1

    wb = Workbook()

    # ---------------- README ----------------
    ws = wb.active
    ws.title = "README"
    lines = [
        ("ACEG Coding Workbook — QM640 Capstone, Group 3", True),
        ("Purpose: double coding of the 223 frozen cases against the ACEG Framework v1 (docs/codebook.md).", False),
        ("", False),
        ("Which cells to edit", True),
        ("Coder1 / Coder2 sheets: only the YELLOW cells (d1–d8, primary_domain, c01–c20, covered, notes). Use the dropdowns.", False),
        ("Gold sheet: only blank YELLOW cells where the coders disagreed; type the agreed value after discussion.", False),
        ("Tracker sheet: coder names in the YELLOW cells. All other cells are formulas or source data — do not edit.", False),
        ("", False),
        ("Coding rules (summary — full rules in docs/codebook.md)", True),
        ("1. Read the case on the Cases sheet (title, description, link). Use only what the advisory says.", False),
        ("2. d1–d8: 1 if the advisory shows that domain's requirement was breached or could be breached; otherwise 0.", False),
        ("3. primary_domain: the single domain (1–8) that most directly enabled the weakness; it must be one coded 1.", False),
        ("4. c01–c20: for controls linked to a domain coded 1, enter 1 = breached (absent, failed or bypassed) or 0 = not breached. "
         "Enter NA for controls whose domain is 0 (those cells turn grey).", False),
        ("5. covered: 1 if at least one breached control would plausibly have prevented the weakness; 0 if the framework has no control for it.", False),
        ("6. n_domains, n_applicable, n_breached, gcs and check are calculated automatically. Fix any message in 'check' before moving on.", False),
        ("7. Code independently. Do not look at the other coder's sheet until the batch is finished.", False),
        ("8. Stay blind to severity: do not look up the CVSS score or severity, and ignore any severity wording in the text when judging controls.", False),
        ("", False),
        ("Schedule", True),
        ("Week 2: both coders code the 20 'W2 pilot' rows; run python src/05_pilot_kappa.py; revise codebook wording where κ < .60.", False),
        ("Weeks 3–5: both coders code batches W3 (68), W4 (68), W5 (67); adjudicate disagreements on the Gold sheet weekly.", False),
        ("End of Week 5: run python src/06_build_coded_dataset.py to create data/processed/cases_coded.csv.", False),
        ("", False),
        ("Example (case CA-018, CVE-2025-54795 — approval-prompt bypass in Claude Code)", True),
    ]
    for i, (t, b) in enumerate(lines, 1):
        c = ws.cell(row=i, column=1, value=t)
        c.font = Font(name=FONT, bold=b, size=12 if i == 1 else 10)
    ex_row = len(lines) + 1
    ex_head = ["d1", "d2", "d3", "d4", "d5", "d6", "d7", "d8", "primary_domain", "c01", "c02", "c03", "c08", "c09",
               "other controls", "covered", "n_applicable", "n_breached", "gcs"]
    ex_val = [1, 0, 0, 1, 0, 0, 0, 0, 1, 0, 1, 0, 1, 0, "NA", 1, 5, 2, 60.0]
    for j, (h, v) in enumerate(zip(ex_head, ex_val), 1):
        ws.cell(row=ex_row, column=j, value=h)
        c = ws.cell(row=ex_row + 1, column=j, value=v)
        c.font = Font(name=FONT, size=10)
        c.border = BORDER
    style_header(ws, len(ex_head), row=ex_row)
    ws.column_dimensions["A"].width = 14
    for j in range(2, len(ex_head) + 1):
        ws.column_dimensions[L(j)].width = 9
    ws.cell(row=ex_row + 3, column=1, value="Note: column A on this sheet holds long text; widen it or read in the formula bar.").font = Font(name=FONT, italic=True, size=9)

    # ---------------- Codebook ----------------
    cb = wb.create_sheet("Codebook")
    cb.append(["Code", "Domain", "Family", "Linked controls"])
    for code, name, fam, cs in DOMAINS:
        cb.append([code, name, fam, ", ".join(f"C{c:02d}" for c in cs)])
    cb.append([])
    cb.append(["Code", "Control", "Domain"])
    ctrl_head = cb.max_row
    for c in range(1, 21):
        cb.append([f"C{c:02d}", CONTROLS[c], DOM_OF[c]])
    cb.append([])
    cb.append(["Families", "A = Agency and oversight (D1, D2, D8); B = Input and code integrity (D4, D6); "
               "C = Data protection (D5); D = Transparency and accountability (D3, D7)"])
    style_header(cb, 4, 1)
    style_header(cb, 3, ctrl_head)
    for col, w in zip("ABCD", (8, 62, 10, 22)):
        cb.column_dimensions[col].width = w
    for row in cb.iter_rows():
        for cell in row:
            if cell.font is None or not cell.font.bold:
                cell.font = Font(name=FONT, size=10, bold=cell.font.bold if cell.font else False)

    # ---------------- Cases (source data, read-only) ----------------
    cs = wb.create_sheet("Cases")
    # Blinding: severity, score, vector, high_impact and user_interaction are NOT shown to coders,
    # because they define the RQ3 outcome and the RQ2 grouping variable.
    ccols = ["case_id", "batch", "product", "component_type", "published", "title", "description", "url"]
    cs.append(ccols)
    for _, r in d.iterrows():
        cs.append([None if (isinstance(r[k], float) and np.isnan(r[k])) else r[k] for k in ccols])
    style_header(cs, len(ccols))
    widths = [9, 9, 16, 13, 11, 40, 90, 40]
    for j, w in enumerate(widths, 1):
        cs.column_dimensions[L(j)].width = w
    for row in cs.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name=FONT, size=9)
            cell.alignment = Alignment(vertical="top", wrap_text=cell.column in (6, 7))
        row[7].hyperlink = row[7].value
        row[7].font = Font(name=FONT, size=9, color="1F4E99", underline="single")
    cs.freeze_panes = "B2"

    # ---------------- Coder sheets ----------------
    v01 = DataValidation(type="list", formula1='"1,0"', allow_blank=True)
    vctrl = DataValidation(type="list", formula1='"1,0,NA"', allow_blank=True)
    vprim = DataValidation(type="list", formula1='"1,2,3,4,5,6,7,8"', allow_blank=True)

    def coder_sheet(name):
        w = wb.create_sheet(name)
        w.append(HEADERS)
        style_header(w, len(HEADERS))
        dv1, dvc, dvp = (DataValidation(type=x.type, formula1=x.formula1, allow_blank=True) for x in (v01, vctrl, vprim))
        for dv in (dv1, dvc, dvp):
            dv.error = "Choose a value from the list"
            dv.showErrorMessage = True
            w.add_data_validation(dv)
        for i in range(2, last + 1):
            w.cell(row=i, column=COL_ID, value=f"=Cases!A{i}")
            w.cell(row=i, column=COL_BATCH, value=f"=Cases!B{i}")
            w.cell(row=i, column=COL_PROD, value=f"=Cases!C{i}")
            w.cell(row=i, column=COL_TITLE, value=f'=IF(Cases!F{i}="",LEFT(Cases!G{i},120),Cases!F{i})')
            dl = [f"{L(COL_D + k)}{i}" for k in range(8)]
            w.cell(row=i, column=COL_NDOM, value=f"=IF(COUNT({dl[0]}:{dl[-1]})=0,\"\",SUM({dl[0]}:{dl[-1]}))")
            napp = "+".join(f"N({dl[k]})*{len(DOMAINS[k][3])}" for k in range(8))
            w.cell(row=i, column=COL_NAPP, value=f"=IF({L(COL_NDOM)}{i}=\"\",\"\",{napp})")
            crange = f"{L(COL_C)}{i}:{L(COL_C + 19)}{i}"
            w.cell(row=i, column=COL_NBR, value=f"=IF({L(COL_NDOM)}{i}=\"\",\"\",COUNTIF({crange},1))")
            w.cell(row=i, column=COL_GCS,
                   value=f"=IF(OR({L(COL_NAPP)}{i}=\"\",{L(COL_NAPP)}{i}=0),\"\",ROUND(100*({L(COL_NAPP)}{i}-{L(COL_NBR)}{i})/{L(COL_NAPP)}{i},1))")
            known = f"COUNT({crange})"
            prim = f"{L(COL_PRIMARY)}{i}"
            w.cell(row=i, column=COL_CHECK, value=(
                f"=IF({L(COL_NDOM)}{i}=\"\",\"\","
                f"IF({L(COL_NDOM)}{i}=0,\"No domain coded 1; \",\"\")"
                f"&IF({prim}=\"\",\"Primary domain missing; \",IF(IFERROR(INDEX({dl[0]}:{dl[-1]},1,{prim}),0)<>1,\"Primary domain not coded 1; \",\"\"))"
                f"&IF({known}<>{L(COL_NAPP)}{i},\"Code 1 or 0 for exactly the controls of domains coded 1 (NA for the rest); \",\"\")"
                f"&IF({L(COL_COVERED)}{i}=\"\",\"Covered missing; \",\"\"))"))
            w.cell(row=i, column=COL_DONE, value=f"=IF(AND({L(COL_NDOM)}{i}<>\"\",{L(COL_CHECK)}{i}=\"\"),1,0)")
            for c in list(range(COL_D, COL_PRIMARY + 1)) + list(range(COL_C, COL_COVERED + 1)) + [COL_NOTES]:
                w.cell(row=i, column=c).fill = YELLOW
            for c in list(range(COL_D, COL_GCS + 1)) + [COL_DONE]:
                w.cell(row=i, column=c).alignment = Alignment(horizontal="center")
            for c in range(1, len(HEADERS) + 1):
                cell = w.cell(row=i, column=c)
                cell.font = Font(name=FONT, size=9)
                cell.border = BORDER
        dv1.add(f"{L(COL_D)}2:{L(COL_D + 7)}{last}")
        dv1.add(f"{L(COL_COVERED)}2:{L(COL_COVERED)}{last}")
        dvp.add(f"{L(COL_PRIMARY)}2:{L(COL_PRIMARY)}{last}")
        dvc.add(f"{L(COL_C)}2:{L(COL_C + 19)}{last}")
        # grey out controls whose domain is coded 0
        for k, (_, _, _, ctrls) in enumerate(DOMAINS):
            dcol = L(COL_D + k)
            for cnum in ctrls:
                col = L(COL_C + cnum - 1)
                w.conditional_formatting.add(f"{col}2:{col}{last}",
                                             FormulaRule(formula=[f"${dcol}2=0"], fill=GREY, stopIfTrue=True))
        w.conditional_formatting.add(f"{L(COL_CHECK)}2:{L(COL_CHECK)}{last}",
                                     FormulaRule(formula=[f"LEN({L(COL_CHECK)}2)>0"], fill=CHECK))
        w.column_dimensions["A"].width = 8
        w.column_dimensions["B"].width = 8
        w.column_dimensions["C"].width = 14
        w.column_dimensions["D"].width = 42
        for c in range(COL_D, COL_GCS + 1):
            w.column_dimensions[L(c)].width = 5.5
        w.column_dimensions[L(COL_PRIMARY)].width = 8
        w.column_dimensions[L(COL_COVERED)].width = 7
        for c in (COL_NDOM, COL_NAPP, COL_NBR, COL_GCS):
            w.column_dimensions[L(c)].width = 8
        w.column_dimensions[L(COL_CHECK)].width = 34
        w.column_dimensions[L(COL_DONE)].width = 6
        w.column_dimensions[L(COL_NOTES)].width = 30
        w.row_dimensions[1].height = 30
        w.freeze_panes = "E2"
        return w

    coder_sheet("Coder1")
    coder_sheet("Coder2")

    # ---------------- Gold (adjudicated) ----------------
    g = wb.create_sheet("Gold")
    g.append(HEADERS[:-2] + ["agree_all", "notes"])
    style_header(g, len(HEADERS))
    coded_cols = list(range(COL_D, COL_PRIMARY + 1)) + list(range(COL_C, COL_COVERED + 1))
    for i in range(2, last + 1):
        for c in (COL_ID, COL_BATCH, COL_PROD, COL_TITLE):
            g.cell(row=i, column=c, value=f"=Coder1!{L(c)}{i}")
        for c in coded_cols:
            a, b = f"Coder1!{L(c)}{i}", f"Coder2!{L(c)}{i}"
            g.cell(row=i, column=c, value=f'=IF(AND({a}&""<>"",{a}&""={b}&""),{a},"")').fill = YELLOW
        dl = [f"{L(COL_D + k)}{i}" for k in range(8)]
        g.cell(row=i, column=COL_NDOM, value=f"=IF(COUNT({dl[0]}:{dl[-1]})=0,\"\",SUM({dl[0]}:{dl[-1]}))")
        napp = "+".join(f"N({dl[k]})*{len(DOMAINS[k][3])}" for k in range(8))
        g.cell(row=i, column=COL_NAPP, value=f"=IF({L(COL_NDOM)}{i}=\"\",\"\",{napp})")
        crange = f"{L(COL_C)}{i}:{L(COL_C + 19)}{i}"
        g.cell(row=i, column=COL_NBR, value=f"=IF({L(COL_NDOM)}{i}=\"\",\"\",COUNTIF({crange},1))")
        g.cell(row=i, column=COL_GCS,
               value=f"=IF(OR({L(COL_NAPP)}{i}=\"\",{L(COL_NAPP)}{i}=0),\"\",ROUND(100*({L(COL_NAPP)}{i}-{L(COL_NBR)}{i})/{L(COL_NAPP)}{i},1))")
        g.cell(row=i, column=COL_CHECK, value=f"=IF(AND(Coder1!{L(COL_DONE)}{i}=1,Coder2!{L(COL_DONE)}{i}=1),"
                                               f"IF(COUNTBLANK({L(COL_D)}{i}:{L(COL_COVERED)}{i})>0,\"Adjudicate blank cells\",\"\"),\"Waiting for both coders\")")
        g.cell(row=i, column=COL_DONE, value=f"=IF(AND(Coder1!{L(COL_DONE)}{i}=1,Coder2!{L(COL_DONE)}{i}=1),"
                                              f"IF(SUMPRODUCT(--(Coder1!{L(COL_D)}{i}:{L(COL_COVERED)}{i}&\"\"<>Coder2!{L(COL_D)}{i}:{L(COL_COVERED)}{i}&\"\"))=0,1,0),\"\")")
        for c in range(1, len(HEADERS) + 1):
            cell = g.cell(row=i, column=c)
            cell.font = Font(name=FONT, size=9)
            cell.border = BORDER
    g.conditional_formatting.add(f"{L(COL_D)}2:{L(COL_COVERED)}{last}",
                                 FormulaRule(formula=[f"AND(${L(COL_CHECK)}2=\"Adjudicate blank cells\",LEN({L(COL_D)}2)=0)"],
                                             fill=CHECK))
    for c, w_ in ((1, 8), (2, 8), (3, 14), (4, 42)):
        g.column_dimensions[L(c)].width = w_
    for c in range(COL_D, COL_GCS + 1):
        g.column_dimensions[L(c)].width = 5.5
    g.column_dimensions[L(COL_CHECK)].width = 24
    g.freeze_panes = "E2"

    # ---------------- Agreement (live κ) ----------------
    a = wb.create_sheet("Agreement")
    a.append(["Live inter-coder agreement (binary domains). κ uses rows where both coders marked the row done."])
    a.append(["Item", "Rows (pilot)", "p_o (pilot)", "κ (pilot)", "Rows (all)", "p_o (all)", "κ (all)", "Flag"])
    style_header(a, 8, 2)
    done1, done2 = f"Coder1!${L(COL_DONE)}$2:${L(COL_DONE)}${last}", f"Coder2!${L(COL_DONE)}$2:${L(COL_DONE)}${last}"
    pil = f"(Cases!$B$2:$B${last}=\"W2 pilot\")"
    for k in range(8):
        col = L(COL_D + k)
        x, y = f"Coder1!${col}$2:${col}${last}", f"Coder2!${col}$2:${col}${last}"
        r = a.max_row + 1
        a.cell(row=r, column=1, value=f"D{k + 1}")
        for start, mask in ((2, f"{pil}*({done1}=1)*({done2}=1)"), (5, f"({done1}=1)*({done2}=1)")):
            nref = f"{L(start)}{r}"
            a.cell(row=r, column=start, value=f"=SUMPRODUCT({mask})")
            po = f"SUMPRODUCT({mask}*({x}={y}))"
            a.cell(row=r, column=start + 1, value=f"=IF({nref}=0,\"\",{po}/{nref})")
            pa, pb = f"(SUMPRODUCT({mask}*({x}=1))/{nref})", f"(SUMPRODUCT({mask}*({y}=1))/{nref})"
            pe = f"({pa}*{pb}+(1-{pa})*(1-{pb}))"
            a.cell(row=r, column=start + 2, value=f"=IF({nref}=0,\"\",IF({pe}=1,1,({L(start + 1)}{r}-{pe})/(1-{pe})))")
        a.cell(row=r, column=8, value=f"=IF(D{r}=\"\",\"\",IF(D{r}<0.6,\"Revise wording (κ < .60)\",\"OK\"))")
    r = a.max_row + 1
    col = L(COL_PRIMARY)
    x, y = f"Coder1!${col}$2:${col}${last}", f"Coder2!${col}$2:${col}${last}"
    a.cell(row=r, column=1, value="Primary domain (exact agreement)")
    for start, mask in ((2, f"{pil}*({done1}=1)*({done2}=1)"), (5, f"({done1}=1)*({done2}=1)")):
        a.cell(row=r, column=start, value=f"=SUMPRODUCT({mask})")
        a.cell(row=r, column=start + 1, value=f"=IF({L(start)}{r}=0,\"\",SUMPRODUCT({mask}*({x}={y}))/{L(start)}{r})")
        a.cell(row=r, column=start + 2, value="see script")
    a.cell(row=r + 2, column=1, value="Multi-category κ for the primary domain, bootstrap CIs and control-level agreement: run python src/05_pilot_kappa.py")
    for row in a.iter_rows(min_row=3):
        for cell in row:
            cell.font = Font(name=FONT, size=10)
            if cell.column in (3, 4, 6, 7):
                cell.number_format = "0.00"
    a["A1"].font = Font(name=FONT, bold=True, size=10)
    a.column_dimensions["A"].width = 30
    for c in "BCDEFG":
        a.column_dimensions[c].width = 12
    a.column_dimensions["H"].width = 24

    # ---------------- Tracker ----------------
    t = wb.create_sheet("Tracker")
    t.append(["Role", "Name"])
    t.append(["Coder 1", ""])
    t.append(["Coder 2", ""])
    t.append(["Adjudicator (if the coders cannot agree)", ""])
    for rr in (2, 3, 4):
        t.cell(row=rr, column=2).fill = YELLOW
    t.append([])
    t.append(["Batch", "Cases", "Week due", "Coder 1 done", "Coder 2 done", "Both done, full agreement", "Gold complete"])
    hdr = t.max_row
    for b_, wk in (("W2 pilot", "Week 2"), ("W3", "Week 3"), ("W4", "Week 4"), ("W5", "Week 5")):
        r = t.max_row + 1
        t.cell(row=r, column=1, value=b_)
        t.cell(row=r, column=2, value=f"=COUNTIF(Cases!$B$2:$B${last},A{r})")
        t.cell(row=r, column=3, value=wk)
        t.cell(row=r, column=4, value=f"=SUMPRODUCT((Coder1!$B$2:$B${last}=A{r})*(Coder1!${L(COL_DONE)}$2:${L(COL_DONE)}${last}=1))")
        t.cell(row=r, column=5, value=f"=SUMPRODUCT((Coder2!$B$2:$B${last}=A{r})*(Coder2!${L(COL_DONE)}$2:${L(COL_DONE)}${last}=1))")
        t.cell(row=r, column=6, value=f"=SUMPRODUCT((Gold!$B$2:$B${last}=A{r})*(Gold!${L(COL_DONE)}$2:${L(COL_DONE)}${last}=1))")
        t.cell(row=r, column=7, value=f"=SUMPRODUCT((Gold!$B$2:$B${last}=A{r})*(Gold!${L(COL_CHECK)}$2:${L(COL_CHECK)}${last}=\"\"))")
    r = t.max_row + 1
    t.cell(row=r, column=1, value="Total")
    for c in range(2, 8):
        if c != 3:
            t.cell(row=r, column=c, value=f"=SUM({L(c)}{hdr + 1}:{L(c)}{r - 1})")
    style_header(t, 2, 1)
    style_header(t, 7, hdr)
    for row in t.iter_rows(min_row=2):
        for cell in row:
            if not (cell.row == hdr):
                cell.font = Font(name=FONT, size=10, bold=(cell.row == r))
    t.column_dimensions["A"].width = 38
    for c in "BCDEFG":
        t.column_dimensions[c].width = 16

    wb.move_sheet("Tracker", offset=-(len(wb.sheetnames) - 2))
    wb.calculation.fullCalcOnLoad = True  # Excel recalculates every formula when the file is opened
    OUT.parent.mkdir(exist_ok=True)
    wb.save(OUT)
    pd.Series(sorted(pilot)).to_csv("coding/pilot_cases.csv", index=False, header=["case_id"])
    print(f"Saved {OUT}; pilot cases: {sorted(pilot)}")


if __name__ == "__main__":
    main()
