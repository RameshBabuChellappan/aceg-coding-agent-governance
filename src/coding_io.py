"""Shared helpers for reading the ACEG coding workbook (coding/ACEG_coding_workbook.xlsx)."""
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

WORKBOOK = Path("coding/ACEG_coding_workbook.xlsx")
DOM = [f"d{i}" for i in range(1, 9)]
CTRL = [f"c{i:02d}" for i in range(1, 21)]
FIELDS = DOM + ["primary_domain"] + CTRL + ["covered"]
LINK = {"d1": [1, 2, 3], "d2": [4, 5], "d3": [6, 7], "d4": [8, 9], "d5": [10, 11],
        "d6": [12, 13, 14], "d7": [15, 16, 17], "d8": [18, 19, 20]}
FAMILY = {1: "A", 2: "A", 8: "A", 4: "B", 6: "B", 5: "C", 3: "D", 7: "D"}
COL = {f: i for i, f in enumerate(FIELDS, start=5)}  # d1 is column E (5) ... covered is column AH (34)


def norm(v):
    """Normalize a typed cell value to int 0/1/2..., 'NA' or None."""
    if v is None or (isinstance(v, str) and v.strip() == "") or (isinstance(v, float) and v != v):
        return None
    if isinstance(v, str):
        s = v.strip().upper()
        if s == "NA":
            return "NA"
        try:
            return int(float(s))
        except ValueError:
            return s
    return int(v)


def read_cases(wb) -> pd.DataFrame:
    ws = wb["Cases"]
    rows = list(ws.iter_rows(min_row=1, max_col=2, values_only=True))
    return pd.DataFrame(rows[1:], columns=["case_id", "batch"])


def read_coder(wb, sheet: str, n: int) -> pd.DataFrame:
    ws = wb[sheet]
    out = []
    for r in range(2, n + 2):
        out.append({f: norm(ws.cell(row=r, column=c).value) for f, c in COL.items()})
    return pd.DataFrame(out, dtype=object)


def load(path: Path = WORKBOOK):
    """Return (cases, coder1, coder2, gold_typed) where gold_typed holds only values typed over formulas."""
    wb_f = load_workbook(path)                  # formulas (to tell typed adjudications from formulas)
    cases = read_cases(wb_f)
    n = len(cases)
    c1, c2 = read_coder(wb_f, "Coder1", n), read_coder(wb_f, "Coder2", n)
    g = wb_f["Gold"]
    typed = []
    for r in range(2, n + 2):
        row = {}
        for f, c in COL.items():
            v = g.cell(row=r, column=c).value
            row[f] = None if (isinstance(v, str) and v.startswith("=")) else norm(v)
        typed.append(row)
    return cases, c1, c2, pd.DataFrame(typed, dtype=object)


def coded(df: pd.DataFrame) -> pd.Series:
    """Rows where a coder has entered all domains, the primary domain and coverage."""
    return df[DOM + ["primary_domain", "covered"]].notna().all(axis=1)


def problems(row: pd.Series) -> list[str]:
    """Consistency checks for one coded row (same rules as the workbook 'check' column)."""
    p = []
    present = [d for d in DOM if row[d] == 1]
    if not present:
        p.append("no domain coded 1")
    prim = row["primary_domain"]
    prim = int(prim) if prim is not None and not (isinstance(prim, float) and prim != prim) else None
    if prim not in range(1, 9) or row[f"d{prim}"] != 1:
        p.append("primary domain not coded 1")
    for d, cs in LINK.items():
        for c in cs:
            v = row[f"c{c:02d}"]
            if row[d] == 1 and v not in (0, 1):
                p.append(f"c{c:02d} needs 1/0")
            if row[d] == 0 and v not in ("NA", None):
                p.append(f"c{c:02d} should be NA")
    if row["covered"] not in (0, 1):
        p.append("covered missing")
    return p


def derive(df: pd.DataFrame, prefix: str = "") -> pd.DataFrame:
    """Add n_domains, n_applicable, n_breached, gcs, domain_family to a frame of codes."""
    out = df.copy()
    d = [prefix + x for x in DOM]
    out[prefix + "n_domains"] = out[d].apply(lambda r: sum(v == 1 for v in r), axis=1)
    app = lambda r: sum(len(LINK[x]) for x in DOM if r[prefix + x] == 1)
    br = lambda r: sum(1 for x in DOM if r[prefix + x] == 1 for c in LINK[x] if r[f"{prefix}c{c:02d}"] == 1)
    out[prefix + "n_applicable"] = out.apply(app, axis=1)
    out[prefix + "n_breached"] = out.apply(br, axis=1)
    out[prefix + "gcs"] = (100 * (out[prefix + "n_applicable"] - out[prefix + "n_breached"])
                           / out[prefix + "n_applicable"].where(out[prefix + "n_applicable"] > 0)).round(1)
    out[prefix + "domain_family"] = out[prefix + "primary_domain"].map(FAMILY)
    return out
